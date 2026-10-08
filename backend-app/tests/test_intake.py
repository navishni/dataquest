import io
import zipfile

import pytest
from openpyxl import Workbook
from pypdf import PdfWriter

from backend.agents import a01_file_validation as a01
from backend.agents import a02_format_router as a02
from backend.agents import a03_native_text as a03
from backend.agents import a08_spreadsheet as a08
from backend.common import audit, config, notify
from backend.common.errors import AgentError, ErrorCode
from backend.common.store import store
from tests.fixtures import make


def up(data, name="a.pdf", password=None):
    return a01.run(a01.FileValidationInput(filename=name, data=data, options=a01.Options(password=password) if password else None))


def events(t):
    return audit.log.db.execute("SELECT count(*) FROM audit WHERE event_type=?", (t,)).fetchone()[0]


# ---------------- 01 ----------------
def test_accepts_pdf_and_persists_encrypted(analyst):
    o = up(make.text_pdf())
    assert o.status == "accepted" and o.detected_mime == "application/pdf" and o.page_count == 1
    assert store.get_file(o.source_id) == make.text_pdf() or store.get_file(o.source_id)[:4] == b"%PDF"
    assert store._files[o.source_id][:3] == b"PF1" and events("file_uploaded") == 1


def test_empty_file_has_specific_error(analyst):
    o = up(b"")
    assert o.status == "rejected" and o.error.code == "EMPTY_FILE"


def test_too_large_uses_config_limit(analyst):
    config.override({"limits": {"max_size_bytes": 100}})
    assert up(make.text_pdf()).error.code == "TOO_LARGE"


def test_extension_mismatch_and_double_extension(analyst):
    assert up(make.text_pdf(), "report.png").error.code == "UNSUPPORTED_FORMAT"
    assert up(b"MZ\x90\x00" + b"\0" * 100, "invoice.pdf.exe").error.reason == "EXECUTABLE"


def test_unknown_content_unsupported(analyst):
    assert up(b"\x00\x01\x02binary", "x.bin").error.code == "UNSUPPORTED_FORMAT"


def test_truncated_pdf_corrupt(analyst):
    assert up(make.text_pdf()[:200]).error.code == "CORRUPT_FILE"


def test_encrypted_pdf_password_flow(analyst):
    w = PdfWriter()
    w.add_blank_page(100, 100)
    w.encrypt("pw123")
    b = io.BytesIO()
    w.write(b)
    assert up(b.getvalue()).error.code == "PASSWORD_REQUIRED"
    assert up(b.getvalue(), password="nope").error.code == "PASSWORD_REQUIRED"
    ok = up(b.getvalue(), password="pw123")
    assert ok.status == "accepted"
    assert "pw123" not in str(audit.log.db.execute("SELECT event_json FROM audit").fetchall())
    assert "pw123" not in repr(store._kv)


def test_macro_office_rejected_and_admin_notified(analyst):
    o = up(make.docx_like({"word/vbaProject.bin": b"x"}), "m.docx")
    assert o.error.reason == "MACRO" and any(i["event_type"] == "file_rejected" for i in notify.notifier.outbox)
    assert events("file_rejected") == 1


def test_zip_bomb_rejected(analyst):
    config.override({"limits": {"max_zip_ratio": 50}})
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("word/document.xml", "<d/>")
        z.writestr("word/big.bin", b"\0" * 20_000_000)
    o = up(b.getvalue(), "bomb.docx")
    assert o.status == "rejected" and o.error.reason == "ZIP_BOMB"


def test_pdf_javascript_rejected(analyst):
    pdf = make.text_pdf().replace(b"%%EOF", b"") + b"\n/JavaScript (app.alert(1))\n%%EOF"
    assert up(pdf).error.reason == "PDF_ACTIVE_CONTENT"


def test_image_bomb_rejected(analyst):
    config.override({"limits": {"max_image_pixels": 1000}})
    assert up(make.png(64, 64), "a.png").error.reason == "IMAGE_BOMB"


def test_duplicate_accepted_with_duplicate_of(analyst):
    data = make.text_pdf()
    a, b = up(data), up(data)
    assert b.status == "accepted" and b.duplicate_of == a.source_id and b.source_id != a.source_id


def test_filename_sanitized(analyst):
    o = up(make.text_pdf(), "../../etc/pass\x00wd‮.pdf")
    assert "/" not in o.sanitized_filename and "\x00" not in o.sanitized_filename and ".." not in o.sanitized_filename


def test_audit_failure_rolls_back_upload(analyst, monkeypatch):
    monkeypatch.setattr(audit.log, "db", None)
    with pytest.raises(AgentError) as e:
        up(make.text_pdf())
    assert e.value.code == ErrorCode.ENGINE_FAILED and not store._files


def test_page_limit(analyst):
    config.override({"limits": {"max_pages": 1}})
    assert up(make.text_pdf(["a", "b"])).error.code == "TOO_LARGE"


# ---------------- 02 ----------------
def test_router_classifies_native_scanned_blank_mixed(analyst):
    import pymupdf
    doc = pymupdf.open()
    p = doc.new_page()
    p.insert_text((72, 100), "A line of native text with plenty of characters in it.")
    doc.new_page()
    p3 = doc.new_page()
    p3.insert_image(p3.rect, stream=make.png(400, 400, "gray"))
    sid = up(doc.tobytes()).source_id
    r = a02.run(a02.RouterInput(source_id=sid))
    assert [u.page_class for u in r.units] == ["native_text", "blank", "scanned"] and r.route == "pdf_mixed"
    assert len({u.unit_id for u in r.units}) == 3
    assert a02.run(a02.RouterInput(source_id=sid)).units == r.units  # stable ids


def test_router_native_only_and_image_routes(analyst):
    assert a02.run(a02.RouterInput(source_id=up(make.text_pdf(["x " * 40])).source_id)).route == "pdf_native"
    assert a02.run(a02.RouterInput(source_id=up(make.png(), "a.png").source_id)).route == "image"


def test_router_xlsx_one_unit_per_sheet(analyst):
    wb = Workbook()
    wb.create_sheet("Two")
    b = io.BytesIO()
    wb.save(b)
    sid = up(b.getvalue(), "w.xlsx").source_id
    r = a02.run(a02.RouterInput(source_id=sid))
    assert r.route == "xlsx" and len(r.units) == 2


def test_router_eml_registers_child_sources(analyst):
    import email.message
    m = email.message.EmailMessage()
    m["From"], m["To"], m["Subject"] = "a@x.com", "b@y.com", "hello"
    m.set_content("body")
    m.add_attachment(make.text_pdf(), maintype="application", subtype="pdf", filename="inv.pdf")
    sid = up(m.as_bytes(), "mail.eml").source_id
    r = a02.run(a02.RouterInput(source_id=sid))
    assert r.route == "eml" and len(r.child_source_ids) == 1 and len(r.units) == 2
    child = store.get("source", r.child_source_ids[0])
    assert child["origin"] == {"type": "email_attachment", "parent_source_id": sid}


def test_router_unknown_source_not_found(analyst):
    with pytest.raises(AgentError) as e:
        a02.run(a02.RouterInput(source_id="nope"))
    assert e.value.code == ErrorCode.NOT_FOUND


def test_router_csv_one_unit(analyst):
    sid = up(b"a,b\n1,2\n", "d.csv").source_id
    assert a02.run(a02.RouterInput(source_id=sid)).units[0].page_number == 1


# ---------------- 03 ----------------
def test_native_text_spans_scaled_to_200dpi(analyst):
    sid = up(make.text_pdf(["Hello World"])).source_id
    o = a03.run(a03.NativeTextInput(source_id=sid, page_number=1))
    s = o.spans[0]
    assert s.text == "Hello World" and o.has_usable_text and s.confidence == 1.0
    assert s.location.page_width == round(595 * 200 / 72) and s.location.bbox[0] == pytest.approx(72 * 200 / 72, abs=3)


def test_garbage_ratio_penalises_private_use_and_replacement_chars():
    assert a03.garbage_ratio("\ue001\ufffdab") == 0.5 and a03.garbage_ratio("clean") == 0.0 and a03.garbage_ratio("") == 0.0


def test_native_text_hidden_text_flagged(analyst):
    import pymupdf
    doc = pymupdf.open()
    p = doc.new_page()
    p.insert_text((72, 100), "visible words", fontsize=11)
    p.insert_text((72, 140), "secret white", fontsize=11, color=(1, 1, 1))
    sid = up(doc.tobytes()).source_id
    o = a03.run(a03.NativeTextInput(source_id=sid, page_number=1))
    assert [s.text for s in o.spans] == ["visible words"] and any(w.code == "HIDDEN_TEXT" for w in o.warnings)


def test_native_text_page_not_found_and_non_pdf(analyst):
    sid = up(make.text_pdf()).source_id
    with pytest.raises(AgentError) as e:
        a03.run(a03.NativeTextInput(source_id=sid, page_number=9))
    assert e.value.code == ErrorCode.NOT_FOUND
    pid = up(make.png(), "a.png").source_id
    with pytest.raises(AgentError) as e2:
        a03.run(a03.NativeTextInput(source_id=pid, page_number=1))
    assert e2.value.code == ErrorCode.UNSUPPORTED_FORMAT


def test_native_text_rotated_page_upright_bbox(analyst):
    import pymupdf
    doc = pymupdf.open()
    p = doc.new_page(width=595, height=842)
    p.insert_text((72, 100), "rotated page text", fontsize=11)
    p.set_rotation(90)
    sid = up(doc.tobytes()).source_id
    o = a03.run(a03.NativeTextInput(source_id=sid, page_number=1))
    s = o.spans[0].location
    assert s.page_width == round(842 * 200 / 72) and s.page_height == round(595 * 200 / 72)
    assert 0 <= s.bbox[0] < s.page_width and 0 <= s.bbox[1] < s.page_height


def test_native_text_deterministic(analyst):
    sid = up(make.text_pdf(["a b c\nd e f"])).source_id
    a = a03.run(a03.NativeTextInput(source_id=sid, page_number=1))
    b = a03.run(a03.NativeTextInput(source_id=sid, page_number=1))
    assert a == b


# ---------------- 08 ----------------
def make_xlsx():
    import datetime
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Item", "Amount", "Rate", "When"])
    ws.append(["Rent", 1234.5, 0.075, datetime.datetime(2025, 4, 30)])
    ws.append(["Total", "=B2*2", "=1/0", None])
    ws["B2"].number_format = '"$"#,##0.00'
    ws["C2"].number_format = "0.00%"
    ws["D2"].number_format = "yyyy-mm-dd"
    ws.merge_cells("A5:B5")
    ws["A5"] = "merged"
    ws.row_dimensions[3].hidden = True
    h = wb.create_sheet("Hidden")
    h.sheet_state = "hidden"
    h["A1"] = "x"
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_spreadsheet_values_formats_formulas(analyst):
    o = a08.run(a08.SpreadsheetInput(source_id=up(make_xlsx(), "w.xlsx").source_id))
    d = {c.ref: c for c in o.sheets[0].cells}
    assert d["B2"].displayed_value == "$1,234.50" and d["C2"].displayed_value == "7.50%" and d["D2"].displayed_value == "2025-04-30"
    assert d["B3"].formula == "=B2*2" and d["B3"].hidden is True and d["A5"].raw_value == "merged"


def test_spreadsheet_hidden_sheets_merged_and_tables(analyst):
    o = a08.run(a08.SpreadsheetInput(source_id=up(make_xlsx(), "w.xlsx").source_id))
    assert [s.hidden for s in o.sheets] == [False, True] and o.sheets[0].merged_ranges == ["A5:B5"]
    assert "A1:D3" in o.sheets[0].probable_tables


def test_spreadsheet_cap_warning_with_true_counts(analyst):
    config.override({"limits": {"max_cells_per_sheet": 3}})
    o = a08.run(a08.SpreadsheetInput(source_id=up(make_xlsx(), "w.xlsx").source_id))
    assert len(o.sheets[0].cells) == 3 and any(w.code == "CELLS_TRUNCATED" and "of" in w.message for w in o.warnings)


def test_spreadsheet_csv_delimiter_sniffing(analyst):
    o = a08.run(a08.SpreadsheetInput(source_id=up(b"a;b;c\n1;2;3\n4;5;6\n", "d.csv").source_id))
    assert {c.ref: c.raw_value for c in o.sheets[0].cells}["C3"] == "6"


def test_spreadsheet_wrong_type_unsupported(analyst):
    with pytest.raises(AgentError) as e:
        a08.run(a08.SpreadsheetInput(source_id=up(make.text_pdf()).source_id))
    assert e.value.code == ErrorCode.UNSUPPORTED_FORMAT


def test_spreadsheet_corrupt_xlsx(analyst):
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w") as z:
        z.writestr("xl/workbook.xml", "<workbook><sheets><sheet name='a'/></sheets></workbook>")
        z.writestr("[Content_Types].xml", "bad")
    sid = up(b.getvalue(), "w.xlsx").source_id
    with pytest.raises(AgentError) as e:
        a08.run(a08.SpreadsheetInput(source_id=sid))
    assert e.value.code == ErrorCode.CORRUPT_FILE


def test_format_value_cases():
    f = a08.format_value
    assert f(0.5, "0%") == "50%" and f(-5, "#,##0;(#,##0)") in ("-5", "(5)") and f(1234567, "#,##0") == "1,234,567" and f(None, None) is None
