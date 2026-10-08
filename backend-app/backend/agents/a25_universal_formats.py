"""Agent 25 - deterministic extraction adapters for non-PDF document formats.

The adapters preserve source-native structure where available (slides, sheets, tables, headings, email headers).
Formats without a source page coordinate system are represented with an explicit unavailable-location reason;
the rendered page is a readable preview and is never presented as the original layout.
"""
import csv
import hashlib
import html
import io
import json
import posixpath
import re
import shutil
import subprocess
import tempfile
import zipfile
from email import policy
from email.parser import BytesParser
from html.parser import HTMLParser
from pathlib import Path, PurePosixPath
from typing import Any
from urllib.parse import unquote, urlsplit
from xml.etree import ElementTree as ET

from pydantic import BaseModel, ConfigDict

from ..common import audit, crypto
from ..common.errors import AgentError, ErrorCode
from ..common.models import Block, Evidence, Location, PageOut, SourceDocument, WarningItem
from ..common.store import store
from .a03_native_text import page_id

PREVIEW_W, PREVIEW_H = 1100, 1500
NO_SOURCE_BOX = "The source format does not store page-image coordinates"
NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
    "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
    "draw": "urn:oasis:names:tc:opendocument:xmlns:drawing:1.0",
}


class UniversalInput(BaseModel):
    model_config = ConfigDict(extra="forbid")
    source_id: str


def _text(value: Any) -> str:
    return "" if value is None else str(value)


def _xml(data: bytes, part: str) -> ET.Element:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            return ET.fromstring(zf.read(part))
    except (zipfile.BadZipFile, KeyError, ET.ParseError, OSError) as exc:
        raise AgentError(ErrorCode.CORRUPT_FILE, f"Document package is corrupt ({part})") from exc


def _zip_parts(data: bytes) -> list[str]:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            return zf.namelist()
    except zipfile.BadZipFile as exc:
        raise AgentError(ErrorCode.CORRUPT_FILE, "Document package is corrupt") from exc


def _xml_text(node: ET.Element, ns: str | None = None) -> str:
    if ns and ns != "text":
        return "".join(n.text or "" for n in node.iter(f"{{{NS[ns]}}}t")).strip()
    if ns == "text":
        parts: list[str] = []

        def visit(current: ET.Element, root: bool = False) -> None:
            local = current.tag.rsplit("}", 1)[-1]
            if local == "s":
                try:
                    count = min(1000, max(1, int(current.get(f"{{{NS['text']}}}c", "1"))))
                except ValueError:
                    count = 1
                parts.append(" " * count)
            elif local == "tab":
                parts.append("\t")
            elif local == "line-break":
                parts.append("\n")
            else:
                if local == "p" and not root and parts and not parts[-1].endswith("\n"):
                    parts.append("\n")
                if current.text:
                    parts.append(current.text)
                for child in current:
                    visit(child)
                    if child.tail:
                        parts.append(child.tail)
                if local == "p" and not root:
                    parts.append("\n")

        visit(node, True)
        return "".join(parts).strip()
    return "".join(node.itertext()).strip()


def _text_item(value: str, kind: str = "paragraph", **extra) -> dict:
    return {"type": kind, "text": value.strip(), "extra": extra}


def _table_item(rows: list[list[dict | str]], header_rows: int = 1, **extra) -> dict:
    cells, n_cols = [], 0
    for row in rows:
        n_cols = max(n_cols, sum(int(c.get("col_span", 1)) if isinstance(c, dict) else 1 for c in row))
    n_rows = len(rows)
    for ri, row in enumerate(rows):
        ci = 0
        for value in row:
            cell = value if isinstance(value, dict) else {"text": value}
            col_span = max(1, int(cell.get("col_span", 1)))
            row_span = max(1, int(cell.get("row_span", 1)))
            raw = _text(cell.get("text", ""))
            normalized = None
            try:
                from .a07_table import parse_numeric
                parsed = parse_numeric(raw)
                if parsed:
                    normalized = {"value": str(parsed[0]), "rule": parsed[1]}
            except Exception:  # noqa: BLE001
                pass
            cells.append({"row": ri, "col": ci, "row_span": row_span, "col_span": col_span,
                          "is_header": bool(cell.get("is_header", ri < header_rows)), "raw_text": raw,
                          "normalized": normalized,
                          "location": {"bbox": None, "coordinate_system": "pixel_top_left", "page_width": PREVIEW_W,
                                       "page_height": PREVIEW_H, "bbox_unavailable_reason": NO_SOURCE_BOX},
                          "confidence": float(cell.get("confidence", 0.98)), "alternatives": [],
                          **({"hidden": True} if cell.get("hidden") else {})})
            ci += col_span
    return {"type": "table", "text": None,
            "data": {"table_id": hashlib.sha256(json.dumps(rows, ensure_ascii=False, default=str).encode()).hexdigest()[:16],
                     "n_rows": n_rows, "n_cols": n_cols, "cells": cells, **extra}}


def _readable_text(data: bytes) -> str:
    from charset_normalizer import from_bytes
    match = from_bytes(data).best()
    if match is not None:
        return str(match)
    return data.decode("utf-8", errors="replace")


def _strip_rtf(data: bytes) -> str:
    src = data.decode("latin-1", errors="replace")
    out, i, depth, skip_depth, uc = [], 0, 0, 0, 1
    skip_groups = {"fonttbl", "colortbl", "stylesheet", "info", "pict", "object", "header", "footer", "*"}
    while i < len(src):
        ch = src[i]
        if ch == "{":
            depth += 1
            if i + 1 < len(src) and src[i + 1] == "\\":
                m = re.match(r"\\([^\W\d]+|\*)", src[i + 1:])
                if m and m.group(1).lower() in skip_groups:
                    skip_depth = depth
            i += 1
            continue
        if ch == "}":
            if skip_depth == depth:
                skip_depth = 0
            depth = max(0, depth - 1)
            i += 1
            continue
        if ch != "\\":
            if not skip_depth and ch not in "\r\n":
                out.append(ch)
            i += 1
            continue
        if i + 1 >= len(src):
            break
        nxt = src[i + 1]
        if nxt in "\\{}":
            if not skip_depth:
                out.append(nxt)
            i += 2
            continue
        if nxt == "'" and i + 3 < len(src):
            try:
                value = bytes.fromhex(src[i + 2:i + 4]).decode("cp1252")
                if not skip_depth:
                    out.append(value)
            except (ValueError, UnicodeDecodeError):
                pass
            i += 4
            continue
        m = re.match(r"\\([a-zA-Z]+)(-?\d+)? ?", src[i:])
        if not m:
            i += 2
            continue
        word, arg = m.group(1).lower(), m.group(2)
        i += len(m.group(0))
        if word == "uc" and arg:
            uc = max(0, int(arg))
        elif word == "u" and arg and not skip_depth:
            code = int(arg) % 65536
            out.append(chr(code))
            i += min(uc, len(src) - i)
        elif not skip_depth and word in ("par", "line"):
            out.append("\n")
        elif not skip_depth and word == "tab":
            out.append("\t")
        elif not skip_depth and word in ("emdash", "endash", "bullet"):
            out.append({"emdash": "—", "endash": "–", "bullet": "•"}[word])
    return re.sub(r"\n{3,}", "\n\n", "".join(out)).strip()


class _HTMLContent(HTMLParser):
    """Collect meaningful HTML blocks and table cells while ignoring navigation and scripts."""
    BLOCKS = {"p", "h1", "h2", "h3", "h4", "h5", "h6", "li", "pre", "blockquote", "figcaption"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.items: list[dict] = []
        self.skip: list[str] = []
        self.current_tag: str | None = None
        self.current: list[str] = []
        self.table_rows: list[list[dict]] | None = None
        self.table_row: list[dict] | None = None
        self.cell: dict | None = None

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        klass = (a.get("class", "") + " " + a.get("id", "")).lower()
        if self.skip:
            if tag in ("script", "style", "nav", "footer", "noscript") or any(k in klass for k in ("cookie", "navbar", "navigation")):
                self.skip.append(tag)
            return
        if tag in ("script", "style", "nav", "footer", "noscript") or any(k in klass for k in ("cookie-banner", "cookie-consent", "navbar", "navigation")):
            self.skip.append(tag)
            return
        if tag == "table":
            self.table_rows = []
        elif tag == "tr" and self.table_rows is not None:
            self.table_row = []
        elif tag in ("td", "th") and self.table_row is not None:
            self.cell = {"text": "", "row_span": int(a.get("rowspan", 1) or 1),
                         "col_span": int(a.get("colspan", 1) or 1), "is_header": tag == "th"}
        elif tag in self.BLOCKS:
            self.current_tag, self.current = tag, []
        elif tag == "br" and self.current_tag:
            self.current.append("\n")
        elif tag == "img":
            alt = a.get("alt") or a.get("title")
            self.items.append(_text_item(alt or "Embedded image", "figure", source="html_image"))

    def handle_data(self, data):
        if self.skip:
            return
        if self.cell is not None:
            self.cell["text"] += data
        elif self.current_tag:
            self.current.append(data)

    def handle_endtag(self, tag):
        if self.skip:
            if self.skip[-1] == tag:
                self.skip.pop()
            return
        if tag in ("td", "th") and self.cell is not None and self.table_row is not None:
            self.cell["text"] = " ".join(self.cell["text"].split())
            self.table_row.append(self.cell)
            self.cell = None
        elif tag == "tr" and self.table_row is not None and self.table_rows is not None:
            self.table_rows.append(self.table_row)
            self.table_row, self.cell = None, None
        elif tag == "table" and self.table_rows is not None:
            if self.table_rows:
                self.items.append(_table_item(self.table_rows, 1, source="html_table"))
            self.table_rows, self.table_row, self.cell = None, None, None
        elif tag == self.current_tag:
            value = " ".join("".join(self.current).split())
            if value:
                kind = "heading" if tag.startswith("h") else "list_item" if tag == "li" else "paragraph"
                self.items.append(_text_item(value, kind, level=int(tag[1]) if tag.startswith("h") else None))
            self.current_tag, self.current = None, []


def _html_items(value: str) -> list[dict]:
    parser = _HTMLContent()
    try:
        parser.feed(value)
        parser.close()
    except Exception:  # noqa: BLE001
        return [_text_item(re.sub(r"<[^>]+>", " ", value), "paragraph")]
    return parser.items


def _plain_items(value: str, markdown: bool = False) -> list[dict]:
    lines = value.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    items, paragraph, table_rows, list_items = [], [], [], []

    def flush_paragraph():
        if paragraph:
            items.append(_text_item(" ".join(x.strip() for x in paragraph).strip(), "paragraph"))
            paragraph.clear()

    def flush_list():
        if list_items:
            items.append(_text_item("\n".join(list_items), "list"))
            list_items.clear()

    def flush_table():
        if table_rows:
            items.append(_table_item(table_rows, 1, source="text_table"))
            table_rows.clear()

    for line in lines:
        stripped = line.strip()
        if not stripped:
            flush_paragraph(); flush_list(); flush_table()
            continue
        if markdown and stripped.startswith("```"):
            flush_paragraph(); flush_list(); flush_table()
            items.append(_text_item(stripped, "code_fence"))
            continue
        hm = re.match(r"^(#{1,6})\s+(.*)$", stripped) if markdown else None
        if hm:
            flush_paragraph(); flush_list(); flush_table()
            items.append(_text_item(hm.group(2), "heading", level=len(hm.group(1))))
            continue
        if markdown and "|" in stripped:
            parts = [x.strip() for x in stripped.strip("|").split("|")]
            if all(re.fullmatch(r":?-{3,}:?", x or "-") for x in parts):
                continue
            flush_paragraph(); flush_list()
            table_rows.append(parts)
            continue
        if re.match(r"^\s*(?:[-*•+]\s+|\d+[.)]\s+)", line):
            flush_paragraph(); flush_table()
            list_items.append(re.sub(r"^\s*(?:[-*•+]\s+|\d+[.)]\s+)", "• ", line))
            continue
        flush_list(); flush_table()
        paragraph.append(stripped)
    flush_paragraph(); flush_list(); flush_table()
    return items


def _json_items(data: bytes) -> list[dict]:
    try:
        obj = json.loads(_readable_text(data))
    except (ValueError, TypeError) as exc:
        raise AgentError(ErrorCode.CORRUPT_FILE, "JSON file is not valid") from exc
    return [{"type": "json_data", "text": json.dumps(obj, ensure_ascii=False, indent=2), "data": {"value": obj}, "extra": {}}]


def _xml_items(data: bytes) -> list[dict]:
    try:
        root = ET.fromstring(data)
    except ET.ParseError as exc:
        raise AgentError(ErrorCode.CORRUPT_FILE, "XML file is not valid") from exc
    lines = [ET.tostring(root, encoding="unicode")]
    return [_text_item("\n".join(lines), "xml_data")]


def _docx_items(data: bytes) -> list[dict]:
    root = _xml(data, "word/document.xml")
    body = root.find(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}body")
    if body is None:
        raise AgentError(ErrorCode.CORRUPT_FILE, "Word document has no body")
    items = []
    w = NS["w"]
    for child in list(body):
        local = child.tag.rsplit("}", 1)[-1]
        if local == "p":
            text = "".join(n.text or "" for n in child.iter(f"{{{w}}}t"))
            text = text.replace("\t", " ").strip()
            math = "".join(n.text or "" for n in child.iter(f"{{{NS['m']}}}t")).strip()
            if math:
                items.append({"type": "equation", "text": math, "data": {"latex": None, "plain_text": math,
                              "crop_url": None, "verified": False}, "extra": {"source": "docx_omml"}})
                if text:
                    items.append(_text_item(text, "paragraph"))
                continue
            if not text:
                continue
            style = child.find("./w:pPr/w:pStyle", {"w": w})
            style_id = style.get(f"{{{w}}}val", "") if style is not None else ""
            kind = "title" if style_id.lower() == "title" else "heading" if "heading" in style_id.lower() else "paragraph"
            items.append(_text_item(text, kind, style=style_id or None))
        elif local == "tbl":
            rows = []
            for tr in child.findall("./w:tr", {"w": w}):
                row = []
                for tc in tr.findall("./w:tc", {"w": w}):
                    value = " ".join(" ".join(_xml_text(p, "w") for p in tc.findall(".//w:p", {"w": w})).split())
                    grid = tc.find("./w:tcPr/w:gridSpan", {"w": w})
                    span = int(grid.get(f"{{{w}}}val", "1")) if grid is not None else 1
                    row.append({"text": value, "col_span": span})
                rows.append(row)
            if rows:
                items.append(_table_item(rows, 2, source="docx_table"))
    if not items:
        raise AgentError(ErrorCode.CORRUPT_FILE, "Word document contains no readable body content")
    return items


def _pptx_chart_items(data: bytes, slide_part: str) -> list[dict]:
    rel_part = str(PurePosixPath(slide_part).parent / "_rels" / (PurePosixPath(slide_part).name + ".rels"))
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            rel_root = ET.fromstring(zf.read(rel_part))
            rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
            chart_targets = [r.get("Target", "") for r in rel_root.findall(f"{{{rel_ns}}}Relationship") if r.get("Type", "").endswith("/chart")]
            parts = zf.namelist()
            items = []
            for target in chart_targets:
                chart_part = str(PurePosixPath(slide_part).parent.joinpath(target))
                stack = []
                norm = []
                for segment in chart_part.split("/"):
                    if segment == "..":
                        if norm: norm.pop()
                    elif segment not in ("", "."):
                        norm.append(segment)
                chart_part = "/".join(norm)
                if chart_part not in parts:
                    continue
                chart = ET.fromstring(zf.read(chart_part))
                c = NS["c"]
                title = " ".join(t.text or "" for t in chart.findall(".//c:title//a:t", NS)).strip() or None
                kind = next((x.tag.rsplit("}", 1)[-1].replace("Chart", "").lower()
                             for x in chart.iter() if x.tag.rsplit("}", 1)[-1].endswith("Chart")), None)
                series = []
                for ser in chart.findall(".//c:ser", NS):
                    name = " ".join(t.text or "" for t in ser.findall("./c:tx//c:v", NS) + ser.findall("./c:tx//a:t", NS)).strip()
                    cats = []
                    for p in ser.findall("./c:cat//c:pt", NS):
                        v = p.find("./c:v", NS)
                        cats.append(v.text if v is not None else "")
                    vals = []
                    for p in ser.findall("./c:val//c:pt", NS):
                        v = p.find("./c:v", NS)
                        try: vals.append(float(v.text))
                        except (AttributeError, ValueError): vals.append(None)
                    points = [{"x": cats[i] if i < len(cats) else i, "y": y}
                              for i, y in enumerate(vals) if y is not None]
                    if points:
                        series.append({"name": name or f"Series {len(series)+1}", "points": points})
                items.append({"type": "chart", "text": title, "data": {"chart_type": kind, "title": title,
                               "series": series, "crop_url": None, "axes": None, "insight_text": None},
                              "extra": {"source": "pptx_chart_xml", "chart_part": chart_part}})
            return items
    except (KeyError, ET.ParseError, zipfile.BadZipFile):
        return []


def _pptx_pages(data: bytes) -> list[list[dict]]:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            names = zf.namelist()
            slides = sorted((n for n in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
                            key=lambda n: int(re.search(r"slide(\d+)", n).group(1)))
            notes = {int(re.search(r"notesSlide(\d+)", n).group(1)): n for n in names
                     if re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", n)}
            result = []
            for i, part in enumerate(slides, 1):
                root = ET.fromstring(zf.read(part))
                tree = root.find(".//p:spTree", NS)
                items = []
                if tree is not None:
                    for shape in list(tree):
                        local = shape.tag.rsplit("}", 1)[-1]
                        if local == "graphicFrame" and shape.find(".//a:tbl", NS) is not None:
                            tab = shape.find(".//a:tbl", NS)
                            rows = []
                            for tr in tab.findall("./a:tr", NS):
                                rows.append([_xml_text(tc, "a") for tc in tr.findall("./a:tc", NS)])
                            if rows: items.append(_table_item(rows, 1, source="pptx_table"))
                        elif local == "graphicFrame" and shape.find(".//c:chart", NS) is not None:
                            items.extend(_pptx_chart_items(data, part))
                        elif local == "pic":
                            items.append(_text_item("Embedded slide image", "figure", source="pptx_picture"))
                        else:
                            for para in shape.findall(".//a:p", NS):
                                value = "".join(t.text or "" for t in para.findall(".//a:t", NS)).strip()
                                if value: items.append(_text_item(value, "paragraph", source="pptx_text"))
                note_part = notes.get(i)
                if note_part:
                    note_root = ET.fromstring(zf.read(note_part))
                    note_text = "\n".join(x.text or "" for x in note_root.findall(".//a:p/a:r/a:t", NS)
                                             if (x.text or "").strip() and not (x.text or "").isdigit())
                    if note_text.strip(): items.append(_text_item(note_text, "speaker_notes"))
                result.append(items)
            if not result:
                raise AgentError(ErrorCode.CORRUPT_FILE, "Presentation contains no readable slides")
            return result
    except (zipfile.BadZipFile, KeyError, ET.ParseError) as exc:
        raise AgentError(ErrorCode.CORRUPT_FILE, "Presentation package is corrupt") from exc


def _odf_pages(data: bytes, mime: str) -> list[list[dict]]:
    root = _xml(data, "content.xml")
    if mime.endswith("presentation"):
        pages = []
        for slide in root.findall(".//draw:page", NS):
            items = []

            def visit(node: ET.Element) -> None:
                local = node.tag.rsplit("}", 1)[-1]
                if local == "table":
                    rows = _odf_table_rows(node)
                    if rows: items.append(_table_item(rows, 1, source="odp_table"))
                    return
                if local in ("p", "h"):
                    value = _xml_text(node, "text")
                    if value: items.append(_text_item(value, "heading" if local == "h" else "paragraph"))
                    return
                for child in node:
                    visit(child)

            visit(slide)
            pages.append(items)
        return pages or [[]]
    if mime.endswith("spreadsheet"):
        pages = []
        for table in root.findall(".//table:table", NS):
            rows = _odf_table_rows(table, repeated=True)
            if rows:
                pages.append([_table_item(rows, 1, sheet=table.get(f"{{{NS['table']}}}name", "Sheet"))])
        return pages or [[]]
    body = root.find(".//{urn:oasis:names:tc:opendocument:xmlns:office:1.0}text")
    if body is None:
        body = root
    items = []

    def visit(node: ET.Element) -> None:
        local = node.tag.rsplit("}", 1)[-1]
        if local == "table":
            rows = _odf_table_rows(node)
            if rows: items.append(_table_item(rows, 1, source="odt_table"))
            return
        if local in ("h", "p"):
            value = _xml_text(node, "text")
            if value: items.append(_text_item(value, "heading" if local == "h" else "paragraph"))
            return
        if local == "list-item":
            value = _xml_text(node, "text")
            if value: items.append(_text_item(value, "list_item"))
            return
        for child in node:
            visit(child)

    visit(body)
    return [items]


def _odf_table_rows(table: ET.Element, repeated: bool = False) -> list[list[str]]:
    rows: list[list[str]] = []
    for tr in table.findall(".//table:table-row", NS):
        row: list[str] = []
        for cell in tr.findall("./table:table-cell", NS) + tr.findall("./table:covered-table-cell", NS):
            value = _xml_text(cell, "text")
            try:
                repeat_cols = min(1000, max(1, int(cell.get(f"{{{NS['table']}}}number-columns-repeated", "1")))) if repeated else 1
            except ValueError:
                repeat_cols = 1
            row.extend([value] * repeat_cols)
        try:
            repeat_rows = min(1000, max(1, int(tr.get(f"{{{NS['table']}}}number-rows-repeated", "1")))) if repeated else 1
        except ValueError:
            repeat_rows = 1
        rows.extend([row] * repeat_rows)
    return rows


def _ods_pages(data: bytes) -> list[list[dict]]:
    root = _xml(data, "content.xml")
    pages = []
    for table in root.findall(".//table:table", NS):
        rows = [[{"text": value, "col_span": 1} for value in row]
                for row in _odf_table_rows(table, repeated=True)]
        rows = [row for row in rows if any(_text(x.get("text")) for x in row)]
        try:
            repeat = min(1000, max(1, int(table.get(f"{{{NS['table']}}}number-rows-repeated", "1"))))
        except ValueError:
            repeat = 1
        if rows:
            page = [_table_item(rows, 1, sheet=table.get(f"{{{NS['table']}}}name", "Sheet"))]
            pages.extend([page] * repeat)
    return pages or [[]]


def _xlsx_pages(source_id: str, data: bytes) -> tuple[list[list[dict]], list[WarningItem]]:
    from openpyxl import load_workbook
    from openpyxl.utils.cell import column_index_from_string, coordinate_from_string, range_boundaries
    from .a08_spreadsheet import SpreadsheetInput, run as spreadsheet_run

    out = spreadsheet_run(SpreadsheetInput(source_id=source_id))
    try:
        wb = load_workbook(io.BytesIO(data), data_only=False, read_only=False)
    except Exception as exc:  # noqa: BLE001
        raise AgentError(ErrorCode.CORRUPT_FILE, "Workbook is corrupt") from exc
    pages = []
    for sheet in out.sheets:
        ws = wb[sheet.name]
        cell_map = {}
        for cell in sheet.cells:
            col_letters, row_num = coordinate_from_string(cell.ref)
            cell_map[(row_num, column_index_from_string(col_letters))] = cell
        ranges = sheet.probable_tables or ([sheet.used_range] if sheet.used_range else [])
        items = []
        for ref in ranges:
            c0, r0, c1, r1 = range_boundaries(ref)
            rows = []
            for ri in range(r0, r1 + 1):
                row = []
                for ci in range(c0, c1 + 1):
                    cell = cell_map.get((ri, ci))
                    value = cell.displayed_value if cell and cell.displayed_value is not None else _text(cell.raw_value if cell else "")
                    row.append({"text": value, "hidden": bool(cell and cell.hidden),
                                "formula": cell.formula if cell else None,
                                "raw_value": cell.raw_value if cell else None})
                rows.append(row)
            items.append(_table_item(rows, 1, source="xlsx_cells", sheet=sheet.name,
                                     hidden_sheet=sheet.hidden, source_range=ref,
                                     merged_ranges=[m for m in sheet.merged_ranges if _range_overlaps(m, ref)]))
        if not items and sheet.name:
            items.append(_text_item(f"{sheet.name} contains no non-empty cells", "note", hidden_sheet=sheet.hidden))
        for chart in ws._charts:
            series = []
            for i, ser in enumerate(chart.ser):
                name = f"Series {i+1}"
                try:
                    if ser.tx and ser.tx.strRef and ser.tx.strRef.strCache and ser.tx.strRef.strCache.pt:
                        name = ser.tx.strRef.strCache.pt[0].v or name
                except AttributeError:
                    pass
                vals = []
                try:
                    if ser.val and ser.val.numRef and ser.val.numRef.numCache:
                        vals = [float(p.v) for p in ser.val.numRef.numCache.pt if p.v is not None]
                except (AttributeError, TypeError, ValueError):
                    vals = []
                cats = []
                try:
                    cat = ser.cat.strRef.strCache if ser.cat and ser.cat.strRef else None
                    if cat: cats = [p.v for p in cat.pt]
                except AttributeError:
                    pass
                points = [{"x": cats[j] if j < len(cats) else j + 1, "y": y} for j, y in enumerate(vals)]
                if points: series.append({"name": name, "points": points})
            title = None
            try:
                title_parts = []
                for para in chart.title.tx.rich.p:
                    runs = getattr(para, "r", []) or []
                    title_parts.extend(getattr(run, "t", "") for run in runs if getattr(run, "t", ""))
                title = " ".join(str(x) for x in title_parts).strip() or None
            except AttributeError:
                pass
            items.append({"type": "chart", "text": title,
                          "data": {"chart_type": type(chart).__name__.replace("Chart", "").lower(), "title": title,
                                   "series": series, "axes": None, "insight_text": None, "crop_url": None},
                          "extra": {"source": "xlsx_chart"}})
        for row in ws.iter_rows():
            for cell in row:
                if cell.comment:
                    items.append(_text_item(f"Comment on {cell.coordinate}: {cell.comment.text}", "cell_comment",
                                            author=cell.comment.author or ""))
        pages.append(items)
    return pages, list(out.warnings)


def _range_overlaps(a: str, b: str) -> bool:
    from openpyxl.utils.cell import range_boundaries
    try:
        ax1, ay1, ax2, ay2 = range_boundaries(a)
        bx1, by1, bx2, by2 = range_boundaries(b)
        return ax1 <= bx2 and bx1 <= ax2 and ay1 <= by2 and by1 <= ay2
    except ValueError:
        return False


def _epub_pages(data: bytes) -> list[list[dict]]:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            container = ET.fromstring(zf.read("META-INF/container.xml"))
            rootfile = next((e.get("full-path") for e in container.iter() if e.tag.rsplit("}", 1)[-1] == "rootfile"), None)
            if not rootfile:
                raise AgentError(ErrorCode.CORRUPT_FILE, "EPUB package has no rootfile")
            opf = ET.fromstring(zf.read(rootfile))
            base = str(PurePosixPath(rootfile).parent)
            manifest = {x.get("id"): x.get("href") for x in opf.iter() if x.tag.rsplit("}", 1)[-1] == "item"}
            pages = []
            for itemref in (x for x in opf.iter() if x.tag.rsplit("}", 1)[-1] == "itemref"):
                href = manifest.get(itemref.get("idref"))
                if not href:
                    continue
                href_path = unquote(urlsplit(href).path)
                part = posixpath.normpath(posixpath.join(base, href_path))
                if part.startswith("../") or part.startswith("/"):
                    raise AgentError(ErrorCode.CORRUPT_FILE, "EPUB manifest points outside the package")
                raw = _readable_text(zf.read(part))
                items = _html_items(raw)
                if items: pages.append(items)
            return pages or [[]]
    except AgentError:
        raise
    except (zipfile.BadZipFile, KeyError, ET.ParseError, OSError) as exc:
        raise AgentError(ErrorCode.CORRUPT_FILE, "EPUB is corrupt") from exc


def _email_items(source_id: str, data: bytes) -> list[dict]:
    try:
        msg = BytesParser(policy=policy.default).parsebytes(data)
    except Exception as exc:  # noqa: BLE001
        raise AgentError(ErrorCode.CORRUPT_FILE, "Email file is malformed") from exc
    items = [{"type": "email_metadata", "text": None,
              "data": {"subject": str(msg.get("subject", "")), "from": str(msg.get("from", "")),
                       "to": str(msg.get("to", "")), "cc": str(msg.get("cc", "")),
                       "date": str(msg.get("date", "")), "message_id": str(msg.get("message-id", "")),
                       "in_reply_to": str(msg.get("in-reply-to", ""))}, "extra": {}}]
    plain, html_body = [], []
    for part in msg.walk():
        if part.is_multipart() or part.get_content_disposition() == "attachment":
            continue
        try:
            value = part.get_content()
        except (LookupError, UnicodeDecodeError, AttributeError):
            continue
        if not isinstance(value, str): continue
        if part.get_content_type() == "text/plain": plain.append(value)
        elif part.get_content_type() == "text/html": html_body.append(value)
    body = "\n\n".join(plain).strip()
    if body:
        items.extend(_plain_items(body))
    elif html_body:
        items.extend(_html_items("\n".join(html_body)))
    for att in sorted(msg.iter_attachments(), key=lambda a: a.get_filename() or ""):
        items.append({"type": "attachment", "text": att.get_filename() or "attachment",
                      "data": {"filename": att.get_filename() or "attachment", "content_type": att.get_content_type(),
                               "size_bytes": len(att.get_payload(decode=True) or b"")}, "extra": {"parent_source_id": source_id}})
    return items


def _spreadsheet_pages(source_id: str, data: bytes) -> tuple[list[list[dict]], list[WarningItem]]:
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
    from .a08_spreadsheet import SpreadsheetInput, run as spreadsheet_run
    out = spreadsheet_run(SpreadsheetInput(source_id=source_id))
    pages = []
    for sheet in out.sheets:
        values = {}
        max_row = max_col = 0
        for cell in sheet.cells:
            col, row = coordinate_from_string(cell.ref)
            col = column_index_from_string(col)
            max_row, max_col = max(max_row, row), max(max_col, col)
            values[(row, col)] = cell
        rows = []
        for r in range(1, max_row + 1):
            row = []
            for c in range(1, max_col + 1):
                cell = values.get((r, c))
                row.append({"text": cell.displayed_value or _text(cell.raw_value) if cell else "",
                            "formula": cell.formula if cell else None,
                            "raw_value": cell.raw_value if cell else None})
            rows.append(row)
        items = []
        if rows:
            items.append(_table_item(rows, 1, source="csv_grid", sheet=sheet.name,
                                     hidden_sheet=sheet.hidden, used_range=sheet.used_range))
        pages.append(items)
    return pages or [[]], list(out.warnings)


def _legacy_xls_pages(data: bytes) -> list[list[dict]]:
    try:
        import xlrd
    except ImportError as exc:
        raise AgentError(ErrorCode.ENGINE_FAILED, "Legacy XLS reading requires xlrd; install backend requirements") from exc
    try:
        book = xlrd.open_workbook(file_contents=data, on_demand=True)
        pages = []
        for sheet in book.sheets():
            rows = [[{"text": _text(sheet.cell_value(r, c))} for c in range(sheet.ncols)] for r in range(sheet.nrows)]
            pages.append([_table_item(rows, 1, sheet=sheet.name, hidden_sheet=sheet.visibility != 0)] if rows else [])
        return pages or [[]]
    except Exception as exc:  # noqa: BLE001
        raise AgentError(ErrorCode.CORRUPT_FILE, "Legacy XLS workbook is corrupt") from exc


def _legacy_office_pages(mime: str, data: bytes) -> list[list[dict]]:
    """Convert legacy Word/PowerPoint through a local LibreOffice install when available."""
    is_word = mime in ("application/vnd.ms-word", "application/msword")
    source_ext, target_ext = ("doc", "docx") if is_word else ("ppt", "pptx")
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if not executable:
        raise AgentError(ErrorCode.UNSUPPORTED_FORMAT,
                         f"Legacy .{source_ext} parsing requires LibreOffice or another compatible converter")
    try:
        with tempfile.TemporaryDirectory(prefix="parsefusion-office-") as tmp:
            work = Path(tmp)
            source = work / f"input.{source_ext}"
            source.write_bytes(data)
            target = work / f"input.{target_ext}"
            profile = work / "profile"
            command = [executable, "--headless", f"-env:UserInstallation={profile.as_uri()}",
                       "--convert-to", target_ext, "--outdir", str(work), str(source)]
            try:
                result = subprocess.run(command, capture_output=True, timeout=45, check=False)
            except subprocess.TimeoutExpired as exc:
                raise AgentError(ErrorCode.TIMEOUT, f"LibreOffice timed out converting legacy .{source_ext}") from exc
            if result.returncode != 0 or not target.is_file():
                raise AgentError(ErrorCode.ENGINE_FAILED,
                                 f"LibreOffice could not convert legacy .{source_ext} to .{target_ext}")
            converted = target.read_bytes()
    except AgentError:
        raise
    except OSError as exc:
        raise AgentError(ErrorCode.ENGINE_FAILED, f"Could not run the LibreOffice .{source_ext} converter") from exc
    return [_docx_items(converted)] if is_word else _pptx_pages(converted)


def _parse(source_id: str, mime: str, data: bytes) -> tuple[list[list[dict]], list[WarningItem]]:
    if mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return [_docx_items(data)], []
    if mime == "application/vnd.openxmlformats-officedocument.presentationml.presentation":
        return _pptx_pages(data), []
    if mime == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        return _xlsx_pages(source_id, data)
    if mime == "application/vnd.oasis.opendocument.text":
        return _odf_pages(data, mime), []
    if mime == "application/vnd.oasis.opendocument.presentation":
        return _odf_pages(data, mime), []
    if mime == "application/vnd.oasis.opendocument.spreadsheet":
        return _ods_pages(data), []
    if mime == "application/epub+zip":
        return _epub_pages(data), []
    if mime in ("application/vnd.ms-word", "application/msword", "application/vnd.ms-powerpoint"):
        return _legacy_office_pages(mime, data), []
    if mime == "application/vnd.ms-excel":
        return _legacy_xls_pages(data), []
    if mime == "text/csv":
        return _spreadsheet_pages(source_id, data)
    if mime == "text/html":
        items = _html_items(_readable_text(data))
        return [items], []
    if mime == "text/rtf":
        return [_plain_items(_strip_rtf(data))], []
    if mime == "text/markdown":
        return [_plain_items(_readable_text(data), markdown=True)], []
    if mime == "text/plain":
        return [_plain_items(_readable_text(data))], []
    if mime == "application/json":
        return [_json_items(data)], []
    if mime == "application/xml":
        return [_xml_items(data)], []
    if mime == "message/rfc822":
        return [_email_items(source_id, data)], []
    raise AgentError(ErrorCode.UNSUPPORTED_FORMAT, f"No universal extraction adapter for {mime}")


def _range_table_text(data: dict) -> list[str]:
    cells = data.get("cells", [])
    rows = {}
    for cell in cells:
        rows.setdefault(cell["row"], {})[cell["col"]] = cell.get("raw_text", "")
    return [" | ".join(row.get(i, "") for i in range(data.get("n_cols", 0))) for _, row in sorted(rows.items())]


def _preview_text(item: dict) -> str:
    if item.get("type") == "table":
        return "\n".join(_range_table_text(item.get("data") or {}))
    if item.get("data", {}).get("value") is not None:
        return _text(item["text"])
    return _text(item.get("text"))


def _render_preview(items: list[dict]) -> bytes:
    from PIL import Image, ImageDraw, ImageFont
    lines = []
    for item in items:
        typ = item.get("type", "paragraph")
        label = {"heading": "", "title": "", "table": "TABLE  ", "chart": "CHART  ",
                 "figure": "FIGURE  ", "equation": "EQUATION  ", "attachment": "ATTACHMENT  ",
                 "speaker_notes": "NOTES  "}.get(typ, "")
        text = _preview_text(item).strip()
        if not text:
            continue
        for line in text.splitlines() or [text]:
            words, current = line.split(), ""
            for word in words:
                if len(current) + len(word) + 1 > 92:
                    lines.append(label + current)
                    label, current = "", word
                else:
                    current = (current + " " + word).strip()
            lines.append(label + current if current else label)
        lines.append("")
    image = Image.new("RGB", (PREVIEW_W, PREVIEW_H), "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("arial.ttf", 18)
        bold = ImageFont.truetype("arialbd.ttf", 21)
    except OSError:
        font = ImageFont.load_default()
        bold = font
    y = 24
    for line in lines:
        if y > PREVIEW_H - 30:
            draw.text((28, PREVIEW_H - 24), "Preview clipped; structured output retains the full extraction.", fill="#555555", font=font)
            break
        is_heading = line.isupper() and line.strip()
        draw.text((28, y), line[:120], fill="#202020", font=bold if is_heading else font)
        y += 28 if is_heading else 23
    out = io.BytesIO()
    image.save(out, format="PNG")
    return out.getvalue()


def _markdown(items_by_page: list[list[dict]]) -> str:
    out = []
    for pn, items in enumerate(items_by_page, 1):
        out.append(f"<!-- page {pn} -->")
        for item in items:
            kind = item.get("type")
            text = _text(item.get("text")).strip()
            if kind == "table":
                data = item.get("data") or {}
                rows = {}
                for c in data.get("cells", []): rows.setdefault(c["row"], {})[c["col"]] = c.get("raw_text", "").replace("|", "\\|")
                grid = [[rows.get(r, {}).get(c, "") for c in range(data.get("n_cols", 0))] for r in range(data.get("n_rows", 0))]
                if grid:
                    out.append("| " + " | ".join(grid[0]) + " |\n|" + "---|" * len(grid[0]))
                    out.extend("| " + " | ".join(row) + " |" for row in grid[1:])
            elif kind in ("heading", "title"):
                level = min(6, max(1, int((item.get("extra") or {}).get("level") or (1 if kind == "title" else 2))))
                out.append("#" * level + " " + text)
            elif kind == "equation":
                out.append("$$\n" + _text((item.get("data") or {}).get("plain_text") or text) + "\n$$")
            elif text:
                out.append(text)
    return "\n\n".join(out)


def run(inp: UniversalInput) -> SourceDocument:
    meta = store.meta(inp.source_id)
    mime = meta.get("detected_mime", "")
    data = store.get_file(inp.source_id)
    page_specs, parse_warnings = _parse(inp.source_id, mime, data)
    page_specs = page_specs or [[]]
    if len(page_specs) > config_max_pages():
        raise AgentError(ErrorCode.TOO_LARGE, "Extracted page count exceeds the configured limit")
    router = store.get("router", inp.source_id) or {}
    units, page_objects = [], []
    from .a02_format_router import unit_id
    for pn, items in enumerate(page_specs, 1):
        pid = page_id(inp.source_id, pn)
        units.append({"unit_id": unit_id(inp.source_id, pn), "page_number": pn,
                      "page_class": "native_text" if items else "blank"})
        blocks = []
        for idx, item in enumerate(items):
            kind = item.get("type", "paragraph")
            text = item.get("text")
            if kind == "equation":
                d = dict(item.get("data") or {})
                d.update({"equation_id": f"{inp.source_id}-eq-{pn}-{idx}", "latex": d.get("latex"),
                          "plain_text": d.get("plain_text") or text, "crop_url": f"/sources/{inp.source_id}/pages/{pn}/image",
                          "verified": False})
            elif kind == "chart":
                d = dict(item.get("data") or {})
                d.update({"chart_id": f"{inp.source_id}-chart-{pn}-{idx}", "crop_url": f"/sources/{inp.source_id}/pages/{pn}/image"})
            elif kind == "figure":
                d = dict(item.get("data") or {})
                d.update({"figure_id": f"{inp.source_id}-figure-{pn}-{idx}", "crop_url": f"/sources/{inp.source_id}/pages/{pn}/image",
                          "caption": text})
            else:
                d = dict(item.get("data") or {})
            bid = hashlib.sha256(f"{inp.source_id}|{pn}|{idx}|{kind}".encode()).hexdigest()[:16]
            loc = Location(bbox=None, page_width=PREVIEW_W, page_height=PREVIEW_H,
                           bbox_unavailable_reason=NO_SOURCE_BOX)
            blocks.append(Block(block_id=bid, type=kind, page_number=pn, reading_order_index=idx,
                                text=text, data=d or None,
                                evidence=Evidence(source_id=inp.source_id, page_id=pid, location=loc,
                                                 extraction_method=mime, confidence=float(item.get("confidence", 0.96))),
                                extra={**(item.get("extra") or {}), **({"needs_review": True} if kind in ("chart", "equation") else {})}))
        image_url = f"/sources/{inp.source_id}/pages/{pn}/image"
        preview = _render_preview(items)
        store.cache_page_image(inp.source_id, pn, preview)
        page = PageOut(page_id=pid, page_number=pn, image_url=image_url, width=PREVIEW_W, height=PREVIEW_H,
                       blocks=blocks)
        page_objects.append(page)
        store.put("layout", pid, {"layout_class": "document", "regions": [], "warnings": []})
        store.put("reading_order", pid, {"ordered_ids": [b.block_id for b in blocks], "reading_order_confidence": 0.95, "warnings": []})
        store.put("consensus", f"{inp.source_id}:{pn}", {"blocks": [], "coverage": {"coverage_score": None, "uncovered_regions": []}})
    warning_list = []
    for raw in parse_warnings:
        warning_list.append(raw if isinstance(raw, WarningItem) else WarningItem.model_validate(raw))
    warning_list.extend(WarningItem(code="FILE_EXTENSION_MISMATCH", message=msg) for msg in (meta.get("ingestion_warnings") or []))
    if not any(page.blocks for page in page_objects):
        warning_list.append(WarningItem(code="NO_CONTENT_EXTRACTED", message="The file opened, but no readable content blocks were produced"))
    router["units"] = units
    router["child_source_ids"] = router.get("child_source_ids", [])
    store.put("router", inp.source_id, router)
    meta["page_count"] = len(page_objects)
    store.put("source", inp.source_id, meta)
    status = "partial" if warning_list else "complete"
    doc = SourceDocument(source_id=inp.source_id, status=status, warnings=warning_list,
                         pages=page_objects, coverage_score=None)
    payload = doc.model_dump(mode="json")
    payload.pop("content_hash", None)
    doc.content_hash = crypto.hash_obj(payload)
    store.put("document", inp.source_id, doc.model_dump(mode="json"))
    store.put("markdown", inp.source_id, _markdown(page_specs))
    audit.append(event_type="universal_format_extracted", object_type="source", object_id=inp.source_id,
                 details={"mime": mime, "pages": len(page_objects), "blocks": sum(len(p.blocks) for p in page_objects), "status": status})
    return doc


def config_max_pages() -> int:
    from ..common import config
    return int(config.get("limits.max_pages", 1000))
